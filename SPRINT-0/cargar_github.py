#!/usr/bin/env python3
"""
Carga G3_Backlog_Sprint0.xlsx en GitHub: épicas, historias de usuario y tareas del Sprint 0
como issues, y las agrega al GitHub Project.

Requisitos:
  - gh CLI instalado y logueado:  gh auth login   (y  gh auth refresh -s project)
  - pip install openpyxl

Uso:
  python cargar_github.py --repo franwe-jpg/gestion-restaurante-G3 --owner franwe-jpg --project 1 --dry-run
  python cargar_github.py --repo franwe-jpg/gestion-restaurante-G3 --owner franwe-jpg --project 1

Completá antes la columna "Usuario GitHub" en la hoja Módulos para que asigne responsables.
Correlo una sola vez: si lo repetís, duplica los issues.
"""
import argparse, json, subprocess, sys
from openpyxl import load_workbook

p = argparse.ArgumentParser()
p.add_argument("--repo", required=True)
p.add_argument("--owner", required=True, help="dueño del Project (usuario u organización)")
p.add_argument("--project", required=True, help="número del Project")
p.add_argument("--xlsx", default="G3_Backlog_Sprint0.xlsx")
p.add_argument("--dry-run", action="store_true")
a = p.parse_args()
_n = 0

def gh(*args):
    global _n
    if a.dry_run:
        _n += 1
        print("gh", " ".join(f'"{x}"' if " " in x else x for x in args))
        return f"https://github.com/{a.repo}/issues/{_n}"
    r = subprocess.run(["gh", *args], capture_output=True, text=True)
    if r.returncode:
        print("ERROR:", r.stderr.strip(), file=sys.stderr)
        return ""
    return r.stdout.strip()

def rows(ws):
    it = ws.iter_rows(values_only=True)
    h = next(it)
    return [dict(zip(h, r)) for r in it if r[0]]

wb = load_workbook(a.xlsx, data_only=True)
mods = rows(wb["Módulos"])
user = {m["Responsable"]: (m["Usuario GitHub"] or "").strip() for m in mods}
nombre = {m["Módulo"]: m["Nombre"] for m in mods}

# Labels
labels = [("épica", "5319E7"), ("historia", "0E8A16"), ("tarea", "FBCA04"),
          ("sprint-0", "1D76DB"), ("transversal", "BFBFBF")]
labels += [(f"{m['Módulo']} {m['Nombre']}", "C5DEF5") for m in mods]
for n, c in labels:
    gh("label", "create", n, "--color", c, "--force", "--repo", a.repo)

def issue(title, body, lbls, assignee=""):
    args = ["issue", "create", "--repo", a.repo, "--title", title, "--body", body]
    for l in lbls: args += ["--label", l]
    if assignee: args += ["--assignee", assignee]
    url = gh(*args)
    if url: gh("project", "item-add", a.project, "--owner", a.owner, "--url", url)
    return url

num = lambda url: "#" + url.rsplit("/", 1)[-1] if url else "?"

# Épicas e historias
epicas = {}
for hu in rows(wb["Story Map"]):
    m = hu["Módulo"]; mlabel = f"{m} {nombre[m]}"
    key = (m, hu["Épica"])
    if key not in epicas:
        epicas[key] = issue(f"[Épica {m}] {hu['Épica']}",
                            f"**Módulo:** {mlabel}\n**Tarea (story map):** {hu['Tarea']}",
                            ["épica", mlabel])
    body = (f"{hu['Historia de usuario']}\n\n**Épica:** {num(epicas[key])}\n"
            f"**Prioridad:** {hu['Prioridad']}\n\n### Criterios de aceptación\n"
            f"- [ ] Dado que ... cuando ... entonces ...\n")
    issue(f"{hu['ID']} {hu['Historia de usuario'].split(', quiero ')[1].split(' para ')[0].capitalize()}",
          body, ["historia", mlabel], user.get(hu["Revisa"], ""))

# Tareas del Sprint 0 (dos pasadas para enlazar dependencias)
tareas = rows(wb["Sprint 0"]); url_de = {}
for t in tareas:
    lbl = ["tarea", "sprint-0"] + (["transversal"] if t["Módulo"] == "Transversal"
                                   else [f"{t['Módulo']} {nombre[t['Módulo']]}"])
    url_de[t["ID"]] = issue(f"{t['ID']} {t['Tarea']}",
                            f"**Entregable:** {t['Entregable']}\n**Semana:** {t['Semana']}",
                            lbl, user.get(t["Responsable"], ""))
for t in tareas:
    deps = [d.strip() for d in str(t["Depende de"]).split(",") if d.strip() not in ("—", "")]
    if deps and url_de.get(t["ID"]):
        dep_txt = ", ".join(f"{d} ({num(url_de.get(d, ''))})" for d in deps)
        gh("issue", "edit", url_de[t["ID"]], "--body",
           f"**Entregable:** {t['Entregable']}\n**Semana:** {t['Semana']}\n**Depende de:** {dep_txt}")

print("Listo." if not a.dry_run else "Fin del dry-run: no se creó nada.")
